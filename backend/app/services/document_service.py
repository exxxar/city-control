import os
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from app.models.problem import Problem
from app.models.document import Document, DocumentType
from app.core.config import settings


class DocumentService:
    def __init__(self, db: Session):
        self.db = db

    def generate_problem_card(self, problem: Problem) -> Document:
        doc = DocxDocument()
        doc.add_heading(f"Обращение № {problem.public_id}", 0)
        doc.add_paragraph(f"Статус: {problem.status.value}")
        doc.add_paragraph(f"Заголовок: {problem.title}")
        doc.add_paragraph(f"Описание: {problem.description or 'Нет описания'}")

        # Сохраняем DOCX
        file_name = f"problem_{problem.public_id}.docx"
        relative_path = f"documents/problems/{file_name}"
        full_path = os.path.join(settings.LOCAL_STORAGE_PATH, relative_path)
        os.makedirs(os.path.dirname(full_path), exist_ok=True)
        doc.save(full_path)

        # Сохраняем запись в БД
        db_doc = Document(
            doc_type=DocumentType.PROBLEM_CARD, file_url=relative_path,
            file_name=file_name, problem_id=problem.id, generated_by_id=problem.created_by_id
        )
        self.db.add(db_doc)
        self.db.commit()
        self.db.refresh(db_doc)
        return db_doc