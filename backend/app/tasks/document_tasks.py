import os
from celery import shared_task
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from docx.shared import Pt

from app.db.session import SessionLocal
from app.models.problem import Problem
from app.models.document import Document, DocumentType
from app.integrations.storage.local_storage import storage
from app.integrations.email.sender import email_sender
from loguru import logger


@shared_task(bind=True, max_retries=3)
def generate_problem_docx(self, problem_id: int):
    """Генерирует DOCX-файл для проблемы и сохраняет его."""
    db: Session = SessionLocal()
    try:
        problem = db.get(Problem, problem_id)
        if not problem:
            logger.warning(f"Problem {problem_id} not found for DOCX generation")
            return

        # 1. Создаем DOCX
        doc = DocxDocument()
        doc.add_heading(f"Обращение № {problem.public_id}", 0)
        doc.add_paragraph(f"Статус: {problem.status.value}").style = 'List Bullet'
        doc.add_paragraph(f"Заголовок: {problem.title}").style = 'List Bullet'
        doc.add_paragraph(f"Описание: {problem.description or 'Нет описания'}")

        # 2. Сохраняем во временный буфер
        file_name = f"problem_{problem.public_id}.docx"
        relative_path = f"documents/problems/{file_name}"

        # python-docx может сохранять в файл или в BytesIO
        import io
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        # 3. Сохраняем в хранилище
        storage.save_file(relative_path, buffer.read())

        # 4. Создаем запись в БД
        db_doc = Document(
            doc_type=DocumentType.PROBLEM_CARD,
            file_url=relative_path,
            file_name=file_name,
            problem_id=problem.id,
            generated_by_id=problem.created_by_id
        )
        db.add(db_doc)
        db.commit()

        logger.info(f"DOCX generated for problem {problem.public_id}")

        # 5. (Опционально) Отправить email заявителю
        # email_sender.send(problem.created_by.email, "Ваше обращение зарегистрировано", ...)

    except Exception as exc:
        logger.error(f"Error generating DOCX for {problem_id}: {exc}")
        self.retry(exc=exc, countdown=60)
    finally:
        db.close()