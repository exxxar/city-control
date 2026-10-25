from io import BytesIO
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from pydantic import BaseModel

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.api.deps import get_current_user
from app.db.session import get_db
from app.models.problem import Problem, ProblemStatus
from app.models.user import User
from app.schemas.problem import ProblemRead, ProblemCreate
from datetime import datetime, date
from enum import Enum

router = APIRouter()


# --- Опционально: Схема для пагинации (можно вынести в app/schemas/problem.py) ---
from pydantic import BaseModel

class ProblemPaginatedResponse(BaseModel):
    items: List[ProblemRead]
    page: int
    per_page: int
    total: int
    pages: int
# ------------------------------------------------------------------------------




def _value(obj, *names, default=None):
    """Возвращает первое существующее непустое поле объекта."""
    if obj is None:
        return default

    for name in names:
        if isinstance(obj, dict):
            value = obj.get(name)
        else:
            value = getattr(obj, name, None)

        if value is not None:
            return value

    return default


def _display(value, default="—"):
    """Преобразует значение в текст для DOCX."""
    if value is None:
        return default

    if isinstance(value, Enum):
        return str(value.value)

    text = str(value).strip()
    return text if text else default


def _format_datetime(value):
    """Форматирует дату или время."""
    if value is None:
        return "—"

    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")

    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")

    return str(value)


def _add_field(document, label, value):
    """Добавляет в документ строку с названием поля и значением."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)

    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(_display(value))

@router.get("", response_model=List[ProblemRead])
def list_problems(
    problem_status: Optional[ProblemStatus] = Query(None, alias="status"),
    db: Session = Depends(get_db),
):
    """Получить список обращений."""
    stmt = select(Problem)

    if problem_status is not None:
        stmt = stmt.where(Problem.status == problem_status)

    problems = db.execute(
        stmt.order_by(Problem.id.desc()).limit(100)
    ).scalars().all()

    return problems


@router.post(
    "",
    response_model=ProblemRead,
    status_code=status.HTTP_201_CREATED,
)
def create_problem(
    payload: ProblemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Создать обращение от имени авторизованного пользователя."""
    problem_data = payload.model_dump(
        exclude_unset=True,
        exclude={"created_by_id"}, # Защита от подмены ID создателя
    )

    problem = Problem(
        **problem_data,
        created_by_id=current_user.id,
    )

    db.add(problem)
    db.commit()
    db.refresh(problem)

    return problem


@router.get("/my", response_model=ProblemPaginatedResponse) # Добавлен response_model
def get_my_problems(
    page: int = Query(1, ge=1),
    per_page: int = Query(10, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Получить обращения текущего пользователя с пагинацией."""
    # Базовое условие фильтрации
    filter_condition = Problem.created_by_id == current_user.id
    
    # ОПТИМИЗИРОВАНО: Прямой подсчет без subquery
    total = db.execute(
        select(func.count(Problem.id)).where(filter_condition)
    ).scalar_one()

    items = db.execute(
        select(Problem)
        .where(filter_condition)
        .order_by(Problem.id.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
    ).scalars().all()

    return {
        "items": items,
        "page": page,
        "per_page": per_page,
        "total": total,
        "pages": (total + per_page - 1) // per_page,
    }


@router.get("/{problem_id}", response_model=ProblemRead)
def get_problem(
    problem_id: int,
    db: Session = Depends(get_db),
):
    """Получить обращение по ID."""
    problem = db.execute(
        select(Problem).where(Problem.id == problem_id)
    ).scalar_one_or_none()

    if problem is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Обращение с ID {problem_id} не найдено",
        )

    return problem


@router.get("/{problem_id}/report.docx", name="download_problem_report")
def download_problem_report(
    problem_id: int,
    db: Session = Depends(get_db),
):
    """Сформировать и скачать DOCX-отчёт по обращению."""
    problem = db.execute(
        select(Problem).where(Problem.id == problem_id)
    ).scalar_one_or_none()

    if problem is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Обращение с ID {problem_id} не найдено",
        )

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("ОТЧЁТ ПО ОБРАЩЕНИЮ ГРАЖДАНИНА")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"№ {problem.id}").bold = True

    document.add_heading("1. Общая информация", level=1)
    status_value = _display(getattr(problem, "status", None))
    status_value = status_value.replace("ProblemStatus.", "")

    _add_field(document, "Номер обращения", problem.id)
    _add_field(document, "Публичный идентификатор", _value(problem, "public_id"))
    _add_field(document, "Дата создания", _format_datetime(_value(problem, "created_at")))
    _add_field(document, "Дата обновления", _format_datetime(_value(problem, "updated_at")))
    _add_field(document, "Статус", status_value)

    priority_id = _value(problem, "priority_id")
    priority_labels = {1: "Критическая", 2: "Высокая", 3: "Средняя", 4: "Низкая"}
    _add_field(document, "Приоритет", priority_labels.get(priority_id, priority_id))
    _add_field(document, "Автор", _value(_value(problem, "created_by"), "full_name", "email"))
    _add_field(document, "Исполнитель", _value(_value(problem, "assigned_to"), "full_name", "email"))
    _add_field(document, "Территория", _value(_value(problem, "territory"), "name", "title"))
    _add_field(document, "Геометка", _value(_value(problem, "geo_tag"), "name", "title"))
    _add_field(document, "Количество подтверждений", _value(problem, "confirmations_count", default=0))
    _add_field(document, "Рейтинг", _value(problem, "rating", default=0))

    document.add_heading("2. Содержание обращения", level=1)
    heading = document.add_paragraph()
    heading.add_run(_display(_value(problem, "title"))).bold = True
    document.add_paragraph(_display(_value(problem, "description"), "Описание не указано"))

    document.add_heading("3. Прикреплённые материалы", level=1)
    attachments = _value(problem, "attachments", default=[]) or []
    if attachments:
        for index, attachment in enumerate(attachments, start=1):
            filename = _value(attachment, "name", "file_name", "original_name", default=f"Вложение {index}")
            file_path = _value(attachment, "url", "file_path", "path")
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(_display(filename)).bold = True
            if file_path:
                paragraph.add_run(f" — {_display(file_path)}")
    else:
        document.add_paragraph("Прикреплённые материалы отсутствуют или не связаны с обращением.")

    document.add_heading("4. История изменений", level=1)
    history = _value(problem, "status_history", "history", default=[]) or []
    if history:
        # Самые старые события идут первыми, если у записей есть дата.
        try:
            history = sorted(history, key=lambda item: _value(item, "created_at", "changed_at") or "")
        except TypeError:
            pass
        for item in history:
            event_date = _format_datetime(_value(item, "created_at", "changed_at"))
            old_status = _value(item, "old_status", "previous_status")
            new_status = _value(item, "new_status", "status")
            comment = _value(item, "comment", "description", "reason")
            parts = [event_date]
            if old_status is not None or new_status is not None:
                parts.append(f"Статус: {_display(old_status, '—')} → {_display(new_status, '—')}")
            if comment:
                parts.append(str(comment))
            document.add_paragraph(" | ".join(parts), style="List Bullet")
    else:
        document.add_paragraph("История изменений отсутствует или не загружена.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Сформировано автоматически системой «Город под контролем»")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="problem-{problem_id}-report.docx"'
            )
        }
    )
